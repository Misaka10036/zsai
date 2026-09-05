"""Deterministic ISO-week resolution and path helpers for Seafile weekly reports."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

logger = logging.getLogger(__name__)


def normalise_path(path: Optional[str]) -> str:
    if not path:
        return "/"
    path = path.strip()
    if not path.startswith("/"):
        path = f"/{path}"
    return path.rstrip("/") or "/"


def parse_mtime(raw_mtime) -> datetime:
    if not raw_mtime:
        return datetime.now(timezone.utc)
    if isinstance(raw_mtime, (int, float)):
        return datetime.fromtimestamp(raw_mtime, tz=timezone.utc)
    if isinstance(raw_mtime, str):
        try:
            return datetime.fromtimestamp(int(raw_mtime), tz=timezone.utc)
        except ValueError:
            pass
        try:
            return datetime.fromisoformat(raw_mtime)
        except ValueError:
            pass
    logger.warning("Unparseable mtime %r, using current time", raw_mtime)
    return datetime.now(timezone.utc)


def host_of(url: str) -> str:
    return (urlparse(url).hostname or "").lower()


def running_in_docker() -> bool:
    return Path("/.dockerenv").exists()


def rewrite_seafile_url_for_runtime(seafile_url: str) -> tuple[str, list[str]]:
    """Rewrite Docker-only hostnames when RAGFlow is running on the host.

    Data-source configs often store ``http://host.docker.internal:8082`` so a
    container can reach Seafile on the host. The host Python process must use
    127.0.0.1 instead — ``host.docker.internal`` resolves to a LAN address
    that commonly returns 502.
    """
    raw = (seafile_url or "").rstrip("/")
    parsed = urlparse(raw)
    hosts: list[str] = []
    hostname = (parsed.hostname or "").lower()
    if hostname:
        hosts.append(hostname)
    if hostname == "host.docker.internal" and not running_in_docker():
        port = parsed.port
        new_host = f"127.0.0.1:{port}" if port else "127.0.0.1"
        rewritten = parsed._replace(netloc=new_host).geturl().rstrip("/")
        hosts.extend(["127.0.0.1", "localhost"])
        logger.info("Rewrote Seafile URL %s -> %s for host runtime", raw, rewritten)
        return rewritten, hosts
    return raw, hosts


def path_in_scope(path: str, sync_path: str) -> bool:
    path = normalise_path(path)
    sync_path = normalise_path(sync_path)
    if sync_path == "/":
        return True
    return path == sync_path or path.startswith(sync_path.rstrip("/") + "/")

WEEK_ID_RE = re.compile(r"(\d{4})-W(\d{2})")
DATE_RE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
FLEX_DATE_RE = re.compile(r"(\d{4})[-./](\d{1,2})[-./](\d{1,2})")


@dataclass(frozen=True)
class WeekWindow:
    week_id: str
    start: datetime
    end: datetime
    previous_week_id: str
    timezone: str


def iso_week_id(value: date) -> str:
    iso = value.isocalendar()
    return f"{iso.year}-W{iso.week:02d}"


def week_bounds(week_id: str, tz: ZoneInfo) -> tuple[datetime, datetime]:
    match = WEEK_ID_RE.fullmatch(week_id)
    if not match:
        raise ValueError(f"invalid week_id {week_id!r}")
    year = int(match.group(1))
    week = int(match.group(2))
    monday = date.fromisocalendar(year, week, 1)
    sunday = monday + timedelta(days=6)
    start = datetime(monday.year, monday.month, monday.day, tzinfo=tz)
    end = datetime(sunday.year, sunday.month, sunday.day, 23, 59, 59, tzinfo=tz)
    return start, end


def previous_iso_week_id(week_id: str) -> str:
    match = WEEK_ID_RE.fullmatch(week_id)
    if not match:
        raise ValueError(f"invalid week_id {week_id!r}")
    year = int(match.group(1))
    week = int(match.group(2))
    monday = date.fromisocalendar(year, week, 1) - timedelta(days=7)
    return iso_week_id(monday)


def parse_week_from_query(query: str) -> str | None:
    if not query:
        return None
    week_match = WEEK_ID_RE.search(query)
    if week_match:
        return f"{week_match.group(1)}-W{week_match.group(2)}"
    date_match = DATE_RE.search(query)
    if date_match:
        day = date(int(date_match.group(1)), int(date_match.group(2)), int(date_match.group(3)))
        return iso_week_id(day)
    return None


def resolve_week(
    now: datetime | None = None,
    timezone_name: str = "Asia/Shanghai",
    week_mode: str = "this_week",
    query: str = "",
) -> WeekWindow:
    try:
        tz = ZoneInfo(timezone_name or "Asia/Shanghai")
    except Exception:
        tz = ZoneInfo("Asia/Shanghai")
        timezone_name = "Asia/Shanghai"

    if now is None:
        now = datetime.now(timezone.utc)
    if now.tzinfo is None:
        now = now.replace(tzinfo=timezone.utc)
    local_now = now.astimezone(tz)
    local_today = local_now.date()

    week_id = parse_week_from_query(query or "")
    mode = (week_mode or "this_week").strip()
    if not week_id:
        if mode == "last_complete_week":
            last_sunday = local_today - timedelta(days=local_today.isoweekday())
            week_id = iso_week_id(last_sunday)
        else:
            week_id = iso_week_id(local_today)

    start, end = week_bounds(week_id, tz)
    if mode != "last_complete_week" and end > local_now:
        end = local_now
    return WeekWindow(
        week_id=week_id,
        start=start,
        end=end,
        previous_week_id=previous_iso_week_id(week_id),
        timezone=timezone_name,
    )


WEEKLY_OUTPUT_FORMATS = ("md", "docx", "pdf")
_OUTPUT_FORMAT_ALIASES = {
    "md": "md",
    "markdown": "md",
    "docx": "docx",
    "doc": "docx",
    "pdf": "pdf",
}
_CONTENT_TYPES = {
    "md": "text/markdown",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def normalize_weekly_output_format(value: str | None) -> str:
    key = (value or "md").strip().lower()
    if key.startswith("."):
        key = key[1:]
    mapped = _OUTPUT_FORMAT_ALIASES.get(key)
    if not mapped:
        raise ValueError(f"output_format must be one of {', '.join(WEEKLY_OUTPUT_FORMATS)}, got {value!r}")
    return mapped


def weekly_report_content_type(output_format: str | None = "md") -> str:
    return _CONTENT_TYPES[normalize_weekly_output_format(output_format)]


def weekly_report_filename(week_id: str, output_format: str | None = "md") -> str:
    if not WEEK_ID_RE.fullmatch(week_id or ""):
        raise ValueError(f"week_id must look like YYYY-Www, got {week_id!r}")
    fmt = normalize_weekly_output_format(output_format)
    return f"weekly-report-{week_id}.{fmt}"


def _markdown_blocks(content: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for raw in (content or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith(("- ", "* ")):
            blocks.append(("li", line[2:].strip()))
        else:
            blocks.append(("p", line))
    return blocks


def _markdown_to_docx(content: str) -> bytes:
    from io import BytesIO

    from docx import Document
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.shared import Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for kind, text in _markdown_blocks(content):
        if kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        elif kind == "h3":
            document.add_heading(text, level=3)
        elif kind == "li":
            document.add_paragraph(text, style="List Bullet")
        else:
            paragraph = document.add_paragraph(text)
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    if not document.paragraphs:
        document.add_paragraph("")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _markdown_to_pdf(content: str) -> bytes:
    from io import BytesIO

    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen.canvas import Canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = BytesIO()
    page = Canvas(buffer, pagesize=A4, pageCompression=0)
    page.setTitle("weekly-report")
    width, height = A4
    y = height - 48
    blocks = _markdown_blocks(content) or [("p", " ")]
    for kind, text in blocks:
        size = 16 if kind == "h1" else 14 if kind in {"h2", "h3"} else 11
        font = "Helvetica" if text.isascii() else "STSong-Light"
        page.setFont(font, size)
        prefix = "• " if kind == "li" else ""
        page.drawString(48, y, f"{prefix}{text}")
        y -= size + 8
        if y < 48:
            page.showPage()
            y = height - 48
    page.save()
    return buffer.getvalue()


def render_weekly_report(content: str, output_format: str | None = "md") -> bytes:
    """Turn synthesizer markdown into the bytes written to the dataset / Seafile."""
    fmt = normalize_weekly_output_format(output_format)
    text = "" if content is None else str(content)
    if fmt == "md":
        return text.encode("utf-8")
    if fmt == "docx":
        return _markdown_to_docx(text)
    return _markdown_to_pdf(text)


def seafile_publish_configured(repo_id: str | None, path: str | None) -> bool:
    """WeeklyPublish writes to Seafile only when both library and path are set."""
    return bool(str(repo_id or "").strip()) and bool(str(path or "").strip())


def resolve_library_id(libraries: list[dict], repo: str) -> str:
    wanted = (repo or "").strip()
    if not wanted:
        raise ValueError("Seafile 资料库不能为空。")
    for lib in libraries or []:
        if (lib.get("id") or "").strip() == wanted:
            return wanted
    for lib in libraries or []:
        if (lib.get("name") or "").strip() == wanted and lib.get("id"):
            return str(lib["id"])
    raise ValueError(f"找不到 Seafile 资料库 {wanted!r}。")


def parse_report_date(value: str | None) -> date | None:
    if not value:
        return None
    match = FLEX_DATE_RE.search(str(value))
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def file_in_week(window: WeekWindow, report_date: str = "", mtime=None) -> bool:
    """Prefer the date encoded in the filename, then fall back to mtime."""
    parsed = parse_report_date(report_date)
    if parsed is not None:
        return window.start.date() <= parsed <= window.end.date()
    modified = parse_mtime(mtime)
    if modified.tzinfo is None:
        modified = modified.replace(tzinfo=timezone.utc)
    start = window.start if window.start.tzinfo else window.start.replace(tzinfo=timezone.utc)
    end = window.end if window.end.tzinfo else window.end.replace(tzinfo=timezone.utc)
    return start <= modified <= end


def search_roots_for_libraries(libraries: list[dict], requested_path: str) -> list[tuple[str, str]]:
    """Map a path onto libraries.

    If the path is a single segment that matches a library name (e.g. ``/日报``
    when a library is named 日报), search that library from ``/``. Otherwise
    search each library from the requested path.
    """
    path = normalise_path(requested_path)
    name = path.strip("/")
    roots: list[tuple[str, str]] = []
    named: list[tuple[str, str]] = []
    for lib in libraries or []:
        lib_id = (lib.get("id") or "").strip()
        if not lib_id:
            continue
        lib_name = (lib.get("name") or "").strip()
        if name and lib_name == name:
            named.append((lib_id, "/"))
        else:
            roots.append((lib_id, path))
    return named or roots
